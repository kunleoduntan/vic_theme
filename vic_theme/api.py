# Copyright (c) 2026, Value Impacts Consulting
# vic_theme/api.py
#
# Generic "export what's on screen" for any List View — unlike a
# per-doctype export (e.g. Pack Profile's own), this takes the doctype,
# the currently applied filters, and the currently visible columns as
# arguments from the client, so one endpoint serves every doctype in
# the instance.

import json
import io
from datetime import datetime

import frappe
from frappe.utils import today, formatdate
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


@frappe.whitelist()
def export_listview_excel(doctype, filters=None, fields=None, accent_color=None):
	"""Export the given doctype's list, with the given filters and fields,
	to a styled .xlsx file. `fields` should be the fieldnames currently
	visible as columns in the caller's List View, so the export mirrors
	what the person is actually looking at rather than every field on
	the doctype. `accent_color` is a hex string (e.g. "#6c5ce7") from the
	caller's active VIC Theme palette, used for the header row fill, so
	the exported file's colors match whichever theme is active in Desk.
	"""
	if isinstance(filters, str):
		filters = json.loads(filters) if filters else {}
	if isinstance(fields, str):
		fields = json.loads(fields) if fields else []
	filters = filters or {}
	fields = fields or []

	if not frappe.has_permission(doctype, "read"):
		frappe.throw(frappe._("Not permitted to read {0}").format(doctype))

	meta = frappe.get_meta(doctype)

	# Always include name; de-duplicate while preserving order
	seen = set()
	ordered_fields = []
	for f in ["name"] + list(fields):
		if f not in seen and meta.has_field(f) or f == "name":
			if f not in seen:
				ordered_fields.append(f)
				seen.add(f)

	if not ordered_fields:
		ordered_fields = ["name"]

	# Column labels from the doctype's own field labels where available
	labels = []
	for fieldname in ordered_fields:
		if fieldname == "name":
			labels.append(frappe._("ID"))
			continue
		df = meta.get_field(fieldname)
		labels.append(df.label if df and df.label else fieldname.replace("_", " ").title())

	records = frappe.get_list(
		doctype,
		filters=filters,
		fields=ordered_fields,
		limit_page_length=0,
		order_by="modified desc",
	)

	# --- Build the workbook -------------------------------------------------
	wb = Workbook()
	ws = wb.active
	ws.title = doctype[:31]  # Excel sheet name limit

	accent = (accent_color or "#6C5CE7").lstrip("#").upper()
	if len(accent) != 6:
		accent = "6C5CE7"

	header_fill = PatternFill(start_color=accent, end_color=accent, fill_type="solid")
	header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
	title_font = Font(name="Calibri", size=14, bold=True, color=accent)
	subtitle_font = Font(name="Calibri", size=10, color="64748B")
	thin_border = Border(
		left=Side(style="thin", color="E2E8F0"),
		right=Side(style="thin", color="E2E8F0"),
		top=Side(style="thin", color="E2E8F0"),
		bottom=Side(style="thin", color="E2E8F0"),
	)
	center_align = Alignment(horizontal="center", vertical="center")

	last_col_letter = get_column_letter(len(ordered_fields))

	# Title + subtitle
	ws.merge_cells(f"A1:{last_col_letter}1")
	ws["A1"] = f"{doctype.upper()} EXPORT"
	ws["A1"].font = title_font
	ws["A1"].alignment = center_align
	ws.row_dimensions[1].height = 28

	ws.merge_cells(f"A2:{last_col_letter}2")
	filter_text = f" | Filters: {json.dumps(filters)}" if filters else ""
	ws["A2"] = f"Generated: {formatdate(today())} | Records: {len(records)}{filter_text}"
	ws["A2"].font = subtitle_font
	ws["A2"].alignment = center_align
	ws.row_dimensions[2].height = 22

	# Header row
	header_row = 4
	for col_idx, label in enumerate(labels, 1):
		cell = ws.cell(row=header_row, column=col_idx, value=label)
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = center_align
		cell.border = thin_border
	ws.row_dimensions[header_row].height = 24

	# Data rows, with a very light zebra tint (mirrors the Desk theme's
	# own subtle-zebra list rows rather than a strong alternating color)
	zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
	for row_offset, rec in enumerate(records, 1):
		row_num = header_row + row_offset
		for col_idx, fieldname in enumerate(ordered_fields, 1):
			value = rec.get(fieldname)
			if value is not None and hasattr(value, "isoformat"):
				value = value.isoformat()
			cell = ws.cell(row=row_num, column=col_idx, value=value)
			cell.border = thin_border
			if row_offset % 2 == 0:
				cell.fill = zebra_fill
		ws.row_dimensions[row_num].height = 20

	# Column widths — a simple heuristic based on label/content length
	for col_idx, fieldname in enumerate(ordered_fields, 1):
		max_len = max(
			[len(str(labels[col_idx - 1]))] + [len(str(r.get(fieldname) or "")) for r in records[:200]]
		)
		ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 4, 12), 45)

	ws.freeze_panes = f"A{header_row + 1}"
	if records:
		ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{header_row + len(records)}"

	output = io.BytesIO()
	wb.save(output)
	output.seek(0)

	timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
	frappe.response["filename"] = f"{doctype.replace(' ', '_')}_{timestamp}.xlsx"
	frappe.response["filecontent"] = output.read()
	frappe.response["type"] = "download"
